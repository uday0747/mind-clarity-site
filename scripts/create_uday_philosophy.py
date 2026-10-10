import json
import os
import subprocess
import shutil
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = str(Path(__file__).resolve().parent.parent / 'references' / 'Uday_Philosophy.docx')
ROOT = Path(__file__).resolve().parent.parent
teaching = json.loads((ROOT / 'teaching-content.json').read_text())
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72); sec.bottom_margin = Inches(0.68)
sec.left_margin = Inches(0.82); sec.right_margin = Inches(0.82)
styles = doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Aptos')
styles['Normal'].font.size=Pt(10.5); styles['Normal'].font.color.rgb=RGBColor(35,35,35)
styles['Normal'].paragraph_format.space_after=Pt(6); styles['Normal'].paragraph_format.line_spacing=1.12
for name,size in [('Title',30),('Heading 1',19),('Heading 2',14),('Heading 3',11.5)]:
    st=styles[name]; st.font.name='Aptos Display' if name in ('Title','Heading 1') else 'Aptos'
    st._element.rPr.rFonts.set(qn('w:eastAsia'),st.font.name); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_before=Pt(15 if name=='Heading 1' else 9); st.paragraph_format.space_after=Pt(5)
    for border in st._element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
quote=styles.add_style('Uday Quote',WD_STYLE_TYPE.PARAGRAPH); quote.base_style=styles['Normal']; quote.font.name='Aptos'; quote.font.size=Pt(11); quote.font.italic=True; quote.font.color.rgb=RGBColor(55,55,55); quote.paragraph_format.left_indent=Inches(.28); quote.paragraph_format.space_before=Pt(6); quote.paragraph_format.space_after=Pt(8)
small=styles.add_style('Small Note',WD_STYLE_TYPE.PARAGRAPH); small.base_style=styles['Normal']; small.font.size=Pt(9); small.font.color.rgb=RGBColor(90,90,90)
for s in doc.sections:
    fp=s.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; fp.style='Small Note'; fp.add_run('Uday Philosophy  •  ')
    run=fp.add_run(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); run._r.addnext(fld)
def p(text='',style=None):
    return doc.add_paragraph(text,style=style)
def h(text,level=1):
    doc.add_heading(text,level=level)
def bullets(items):
    for x in items: doc.add_paragraph(x,style='List Bullet')
def label(text):
    q=doc.add_paragraph(style='Small Note'); q.add_run(text.upper()).bold=True


CONTENT = ROOT / 'source-notes/philosophy-consolidated-content.md'
node = os.environ.get('MIND_CLARITY_NODE') or shutil.which('node')
if not node:
    raise RuntimeError('Node.js is required to read the current website FAQ source')
code = "import {questions,enquiryGroups} from './questions.js'; console.log(JSON.stringify([{title:'Core questions',entries:questions},...enquiryGroups]));"
groups = json.loads(subprocess.check_output([node, '--input-type=module', '-e', code], cwd=ROOT, text=True))

p('MIND CLARITY  /  PHILOSOPHY REFERENCE', 'Small Note')
doc.add_paragraph('Uday Philosophy', 'Title')
p('The Creator and the game of experiences')
p('Uday Kumar  |  Updated 10 October 2026  |  Version 1.8', 'Small Note')
p('The current teaching, later clarifications, traditional foundations, everyday examples and Beautiful World vision are organised by topic. Original statements remain in separate source files; this document is the consolidated reference for sharing and continuing the work.')
p('Claims stated as Uday\'s views preserve his position. Editorial distinctions separate those views from verified findings and historical assistant proposals.', 'Small Note')
doc.add_page_break()
h('Contents')
text = CONTENT.read_text()
for line in text.splitlines():
    if line.startswith('## '):
        p(line[3:])
p('12 Question catalogue')
doc.add_page_break()
seen_body = set()
def unique_p(text, style=None):
    normalized = ' '.join(text.split())
    if normalized in seen_body:
        return
    seen_body.add(normalized)
    p(text, style)

def question_section(question, group_title):
    q = question.lower()
    if 'guilt' in q:
        return '5 and 7'
    if 'guarantee perfect health or wealth' in q:
        return '4 and 5'
    if 'artificial intelligence' in q or 'fear it' in q:
        return 8
    if any(word in q for word in ['health', 'oxygen', 'meditation', 'astrology', 'planets', 'eeg']):
        return 4
    if any(word in q for word in ['supernatural', 'past lives', 'painful experiences', 'everything that will happen', 'effort', 'game cannot', 'surrender the same']):
        return 2
    if 'movie' in q or 'sadness and compassion' in q or 'steadiness' in q:
        return 6 if 'practical needs' not in q else 7
    if any(word in q for word in ['teacher', 'deadline', 'write down', 'affirmation', 'question the', 'begin', 'repeating a belief', 'prove we created', 'disagree', 'agree with this']):
        return 3
    if any(word in q for word in ['central statement', 'refer to', 'creator enter']):
        return 1
    if any(word in q for word in ['confidence', 'personality', 'overthinking', 'guilt', 'depression', 'possessions', 'relationships fit']):
        return 5
    if 'bhakti' in q:
        return 4
    if 'website trying' in q:
        return 10
    if 'uday or this website make mistakes' in q:
        return 3
    if 'creator mean accepting harm' in q or group_title in ['Care, responsibility, and everyday life','Decisions, steadiness, and learning','Family, leadership, and honest enquiry']:
        return 7
    if 'mind does not become silent' in q:
        return 4
    if 'why is responsibility' in q:
        return 1
    if 'sadness' in q or 'emotional pain' in q:
        return 6
    return 7

for block in text.splitlines():
    block = block.strip()
    if not block or block.startswith('# '):
        continue
    if block.startswith('## '):
        h(block[3:],1)
    elif block == '[ORIGINAL_MULA]':
        label('Original statement supplied by Uday on 16 September 2026')
        unique_p(teaching['original'], 'Uday Quote')
    elif block == '[SCRIPTURAL_POINTERS]':
        for item in [
            'Prajñānam Brahma — “Consciousness is Brahman.” (Aitareya Upanishad 3.1.3; associated with the Rig Veda)',
            'Aham Brahmāsmi — “I am Brahman.” (Brihadaranyaka Upanishad 1.4.10; associated with the Yajur Veda)',
            'Tat Tvam Asi — “That thou art” / “You are that ultimate reality.” (Chandogya Upanishad 6.8.7; associated with the Sama Veda)',
            'Ayam Ātmā Brahma — “This Self is Brahman.” (Mandukya Upanishad 2; associated with the Atharva Veda)',
            'So’ham Asmi — “That person there—I am he.” (Isha Upanishad 16; Shukla Yajur Veda 40.16)',
            'Aham Rudrebhir Vasubhiś Carāmi — “I move with the Rudras and the Vasus,” expressing the all-pervading divine power speaking through the Devi Sukta. (Rig Veda 10.125.1)',
            'Sarvam Khalvidam Brahma — “All this is indeed Brahman.” (Chandogya Upanishad 3.14.1; associated with the Sama Veda)',
            'Aham Manur Abhavam Sūryaś Ca — “I became Manu; I became the Sun,” expressing an expansive realization of identity with cosmic life. (Rig Veda 4.26.1)',
            'Yad Eveha Tad Amutra — “What is here is there; what is there is here.” (Katha Upanishad 2.1.10; associated with the Yajur Veda)',
            'Sa Eṣo’ṇimā Etad Ātmyam Idaṁ Sarvam — “That subtle essence is the Self of all this.” (Chandogya Upanishad 6.8.7; associated with the Sama Veda)'
        ]: p(item, 'List Bullet')
    elif block == '[PRACTICES_AND_EXAMPLES]':
        h('Five optional reflection steps', 2)
        for title, description in teaching['steps']:
            h(title,3)
            unique_p(description)
        for example in teaching['cases']:
            h(example['title'],2)
            for key in ['description','connection','example','scene','remember','participation','reflection']:
                if example.get(key):
                    unique_p(example[key])
    elif block == '[FAQ_CATALOGUE]':
        h('12 Question catalogue')
        unique_p('All 56 current website questions are listed here with the section containing their consolidated explanation. The full editorial website answers remain in questions.js; they are not repeated as a second account of the philosophy in this document.')
        seen_questions = set()
        for group in groups:
            h(group['title'], 2)
            for question, answer in group['entries']:
                if question in seen_questions:
                    continue
                seen_questions.add(question)
                unique_p(f"{question}  See section {question_section(question, group['title'])}.")
    else:
        unique_p(block)

h('Source locations and comparison record',2)
for entry in [
    'references/archive/Uday_Beautiful_World_Guide.docx — original guide of 16 September 2026',
    'references/archive/Uday_Philosophy_before_self_guided_design.docx — historical reference of 18 September 2026',
    'references/archive/enlightenment-mission-statement-2026-10-01.md — earlier 19-part statement collection',
    'source-notes/project-original-philosophy-statements.json — selected substantive text from five earlier project chats',
    'source-notes/2026-10-10-uday-original-statements-te.md — original Telugu statements and confirmed correction',
    'teaching-content.json — shared Mula Siddhanta, practices and six examples',
    'questions.js — 56 current website question and answer pairs',
    'references/archive/decision-making-evidence-review-2026-10-02.md and references/archive/inner-clarity-and-eeg-review-2026-10-02.md — retained evidence reviews',
    'references/archive/beautiful-world-review.md — historical guide review',
    'audits/philosophy-chat-comparison-2026-10-10.md — reviewed chats and reconciliation decisions'
]:
    unique_p(entry, 'Small Note')
doc.core_properties.title = 'Uday Philosophy'
doc.core_properties.author = 'Uday Kumar'
doc.core_properties.subject = 'The Creator and the game of experiences'
doc.save(OUT)
print(f'{OUT}\nConsolidated reference version 1.8 with {len(seen_questions)} FAQ catalogue entries')
